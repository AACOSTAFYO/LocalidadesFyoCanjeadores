"""
Script to create a sample Excel file with example data.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill


def create_sample_excel(filename='sample_data.xlsx'):
    """
    Create a sample Excel file with example data for lists a and b.
    
    Args:
        filename: Name of the file to create
    """
    # Create a new workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    
    # Add headers with formatting
    ws['A1'] = 'Lista A'
    ws['B1'] = 'Lista B'
    
    # Style headers
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    ws['A1'].font = header_font
    ws['B1'].font = header_font
    ws['A1'].fill = header_fill
    ws['B1'].fill = header_fill
    
    # Add sample data for Lista A
    sample_list_a = [
        'Buenos Aires',
        'Córdoba',
        'Rosario',
        'Mendoza',
        'La Plata',
        'San Miguel de Tucumán',
        'Mar del Plata',
        'Salta',
        'Santa Fe',
        'San Juan'
    ]
    
    # Add sample data for Lista B
    sample_list_b = [
        'Canjeador 001',
        'Canjeador 002',
        'Canjeador 003',
        'Canjeador 004',
        'Canjeador 005',
        'Canjeador 006',
        'Canjeador 007',
        'Canjeador 008',
        'Canjeador 009',
        'Canjeador 010'
    ]
    
    # Write data to cells
    for idx, value in enumerate(sample_list_a, start=2):
        ws[f'A{idx}'] = value
    
    for idx, value in enumerate(sample_list_b, start=2):
        ws[f'B{idx}'] = value
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 25
    
    # Save the workbook
    wb.save(filename)
    print(f"Sample Excel file created: {filename}")
    print(f"  - Lista A: {len(sample_list_a)} items")
    print(f"  - Lista B: {len(sample_list_b)} items")


if __name__ == "__main__":
    create_sample_excel()
